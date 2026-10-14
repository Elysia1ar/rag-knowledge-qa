"""基于 LangChain 的多格式文档加载封装。"""
from pathlib import Path

# 1.x 中 TextLoader 仍在 community，但 PyPDFLoader 等可能有变化
# 建议优先用 langchain_community 的 loader，如果 community 0.4.2 不支持
# 可以改用 pypdf 直接读取（我们后续可以封装自己的 loader）
from langchain_community.document_loaders import (
    PyPDFLoader,
    TextLoader,
    UnstructuredWordDocumentLoader,
)
from langchain_core.documents import Document

from app.core.exceptions import DocumentLoadError, UnsupportedFileTypeError
from app.core.logging import get_logger
from app.utils.file_utils import SUPPORTED_EXTENSIONS


logger = get_logger(__name__)

class _DocxStyleLoader:
    """自定义 docx loader：用 python-docx 读样式，输出带 style_level 的 Document 列表。

    接口与 LangChain 的 BaseLoader 兼容（构造传路径，调用 .load()）。
    """

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)

    def load(self) -> list[Document]:
        from docx import Document as DocxDocument

        d = DocxDocument(str(self.path))
        docs: list[Document] = []

        for i, para in enumerate(d.paragraphs):
            text = (para.text or "").strip()
            if not text:
                continue

            style_name = (para.style.name if para.style else "") or ""

            # 从样式名解析层级
            level = 0
            if style_name.startswith("Heading "):
                try:
                    level = int(style_name.split()[-1])
                except ValueError:
                    level = 1
            elif style_name == "Title":
                level = 1

            category = "Title" if level > 0 else "NarrativeText"

            docs.append(Document(
                page_content=text,
                metadata={
                    "source": str(self.path),
                    "file_name": self.path.name,
                    "file_type": "docx",
                    "category": category,
                    "style_name": style_name,
                    "style_level": level,     # 0=正文, 1/2/3=标题层级
                    "page": i,
                },
            ))
        return docs

_LOADER_MAP = {
    ".pdf": PyPDFLoader,
    ".txt": lambda p: TextLoader(p, encoding="utf-8"),
    ".md": lambda p: TextLoader(p, encoding="utf-8"),   # ← 改成 TextLoader
    ".docx": _DocxStyleLoader,
}


def load_single_file(path: str | Path) -> list[Document]:
    """按扩展名派发到对应 loader，并补全元数据。"""
    p = Path(path)
    ext = p.suffix.lower()

    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(f"不支持的文件类型: {ext}")

    loader_cls = _LOADER_MAP.get(ext)
    if loader_cls is None:
        raise UnsupportedFileTypeError(f"未注册 loader: {ext}")

    try:
        loader = loader_cls(str(p))
        docs = loader.load()
    except Exception as e:
        raise DocumentLoadError(f"加载失败: {p.name}", detail=str(e)) from e

    for i, doc in enumerate(docs):
        doc.metadata.setdefault("source", str(p))
        doc.metadata.setdefault("file_name", p.name)
        doc.metadata.setdefault("file_type", ext.lstrip("."))
        doc.metadata.setdefault("page", doc.metadata.get("page", i))
    logger.info("加载 %s，共 %d 段", p.name, len(docs))
    return docs


def load_directory(directory: str | Path | None = None) -> list[Document]:
    """加载目录下所有支持格式的文件。"""
    from app.core.config import get_settings
    directory = Path(directory) if directory else get_settings().raw_dir

    all_docs: list[Document] = []
    for path in sorted(directory.rglob("*")):
        if not path.is_file():
            continue
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        try:
            all_docs.extend(load_single_file(path))
        except DocumentLoadError as e:
            logger.warning("跳过 %s: %s", path.name, e)

    logger.info("目录 %s 共加载 %d 段", directory, len(all_docs))
    return all_docs
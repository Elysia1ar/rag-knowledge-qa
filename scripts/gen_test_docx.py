"""生成测试用多级标题 Word 文件。"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def main() -> None:
    from docx import Document

    doc = Document()

    # 一级标题
    doc.add_heading("RAG 系统详解", level=1)
    doc.add_paragraph(
        "RAG（Retrieval-Augmented Generation）是一种将检索与生成结合的技术架构。"
        "它在生成回答前先从外部知识库检索相关内容，让模型基于真实材料作答。"
    )

    # 二级标题
    doc.add_heading("一、核心组件", level=2)
    doc.add_paragraph(
        "RAG 系统由四大核心组件构成：文档加载器、Embedding 模型、"
        "向量数据库和重排序模型。每个组件都有多种实现可选。"
    )

    # 三级标题
    doc.add_heading("1.1 文档加载器", level=3)
    doc.add_paragraph(
        "文档加载器负责把各种格式的文件转换为统一的 Document 对象。"
        "常见格式包括 PDF、Word、Markdown、TXT 等。"
        "不同格式需要不同的解析策略，比如 PDF 按页、Markdown 按标题。"
    )

    doc.add_heading("1.2 Embedding 模型", level=3)
    doc.add_paragraph(
        "Embedding 模型把文本编码为稠密向量。中文场景推荐 BGE 系列："
        "bge-small-zh-v1.5 向量维度 512，适合快速验证；"
        "bge-large-zh-v1.5 向量维度 1024，检索质量更高。"
    )

    # 另一个二级标题
    doc.add_heading("二、检索流程", level=2)
    doc.add_paragraph(
        "检索流程包括向量召回、BM25 召回、RRF 融合、Cross-Encoder 重排四个阶段。"
    )

    doc.add_heading("2.1 混合检索", level=3)
    doc.add_paragraph(
        "向量检索擅长语义匹配，BM25 擅长关键词精确匹配。"
        "两者互补，通过 RRF 算法融合后效果显著优于单一策略。"
    )

    doc.add_heading("2.2 重排序", level=3)
    doc.add_paragraph(
        "重排序使用 Cross-Encoder 对候选片段精细打分。"
        "相比 Bi-Encoder 的双塔结构，Cross-Encoder 通过 query-document 交互"
        "获得更高精度，但推理速度慢，通常只对 Top-K 候选使用。"
    )

    out = Path("data/raw/docx_test.docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"✅ 生成 {out}")


if __name__ == "__main__":
    main()